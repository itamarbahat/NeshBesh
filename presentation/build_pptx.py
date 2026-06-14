# -*- coding: utf-8 -*-
"""Generate NeshBesh class presentation (Hebrew, RTL)."""
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU = 914400

# ---- Palette -------------------------------------------------------------
BG      = RGBColor(0x0F, 0x1B, 0x2D)   # deep navy
PANEL   = RGBColor(0x16, 0x26, 0x3D)   # panel
BLUE    = RGBColor(0x35, 0x9B, 0xF0)   # Nesh
GREEN   = RGBColor(0x2E, 0xCC, 0x71)   # Besh
WHITE   = RGBColor(0xF2, 0xF5, 0xFA)
MUTED   = RGBColor(0xA9, 0xB8, 0xCC)
GOLD    = RGBColor(0xF1, 0xC4, 0x0F)
DARKTRI = RGBColor(0x1D, 0x2F, 0x4A)

FONT = "Arial"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    r._element.addprevious(r._element)  # keep at back (already first)
    return s


def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set('rtl', '1')


def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return tb, tf


def style(run, size, color, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic; run.font.name = font


def para(tf, first=False, align=PP_ALIGN.RIGHT, rtl=True, space_after=8, space_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if rtl:
        set_rtl(p)
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    return p


def rect(s, x, y, w, h, color, line=None, radius=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = line; shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def accent_header(s, title_he, num):
    """Title bar with green/blue accent stripe + Hebrew title + slide number."""
    # accent stripe on the right (RTL)
    rect(s, SW - Inches(0.55), Inches(0.55), Inches(0.16), Inches(0.95), BLUE)
    tb, tf = box(s, Inches(0.6), Inches(0.5), Inches(11.9), Inches(1.05),
                 anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, first=True)
    style(p.add_run(), 1, WHITE)  # placeholder removed below
    p.runs[0].text = title_he
    style(p.runs[0], 34, WHITE, bold=True)
    # underline accent line
    rect(s, Inches(0.6), Inches(1.6), Inches(12.1), Pt(2.2), GREEN)
    # slide number chip
    tb2, tf2 = box(s, Inches(0.4), SH - Inches(0.55), Inches(1.0), Inches(0.4))
    p2 = para(tf2, first=True, align=PP_ALIGN.LEFT, rtl=False)
    style(p2.add_run(), 12, MUTED)
    p2.runs[0].text = f"{num} / 8"


def bullet(tf, lead, rest="", first=False, size=18, lead_color=GREEN,
           space_after=10):
    p = para(tf, first=first, space_after=space_after)
    r1 = p.add_run(); r1.text = "●  "
    style(r1, size, lead_color, bold=True)
    r2 = p.add_run(); r2.text = lead
    style(r2, size, WHITE, bold=True)
    if rest:
        r3 = p.add_run(); r3.text = " — " + rest
        style(r3, size, MUTED)
    return p


# =========================================================================
# SLIDE 1 — COVER
# =========================================================================
s = slide()

# Title NeshBesh (Nesh blue, Besh green)
tb, tf = box(s, Inches(0), Inches(0.55), SW, Inches(1.5), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER, rtl=False)
r = p.add_run(); r.text = "Nesh"; style(r, 70, BLUE, bold=True)
r = p.add_run(); r.text = "Besh"; style(r, 70, GREEN, bold=True)
# subtitle
tb, tf = box(s, Inches(0), Inches(1.85), SW, Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER)
r = p.add_run(); r.text = "גרסה דיגיטלית למשחק השש־בש המשפחתי"
style(r, 20, MUTED)

# ---- Backgammon board ----
bL = Inches(2.85); bT = Inches(2.7); bW = Inches(7.63); bH = Inches(4.0)
frame = rect(s, bL, bT, bW, bH, RGBColor(0x0A, 0x14, 0x22),
             line=GOLD, radius=True)
frame.line.width = Pt(3)
pad = Inches(0.18)
pL, pT = bL + pad, bT + pad
pW, pH = bW - 2*pad, bH - 2*pad
play_bg = rect(s, pL, pT, pW, pH, RGBColor(0x12, 0x20, 0x36))
bar_w = Inches(0.42)
bar_x = pL + (pW - bar_w)//2
rect(s, bar_x, pT, bar_w, pH, RGBColor(0x0A, 0x14, 0x22))
half_w = (pW - bar_w)//2
tri_base = half_w // 6
tri_h = int(pH * 0.42)


def triangle(pts, color):
    fb = s.shapes.build_freeform(Emu(pts[0][0]), Emu(pts[0][1]), scale=1)
    fb.add_line_segments([(Emu(pts[1][0]), Emu(pts[1][1])),
                          (Emu(pts[2][0]), Emu(pts[2][1]))], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.color.rgb = RGBColor(0x0A, 0x14, 0x22); shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def quad_x(i):
    # i = 0..11 across both halves accounting for the bar gap
    if i < 6:
        return pL + i*tri_base
    return bar_x + bar_w + (i-6)*tri_base


topY = int(pT)
botY = int(pT + pH)
for i in range(12):
    x0 = int(quad_x(i)); x1 = x0 + int(tri_base)
    col = BLUE if i % 2 == 0 else GREEN
    # top, pointing down
    triangle([(x0, topY), (x1, topY), ((x0+x1)//2, topY+tri_h)], col)
    # bottom, pointing up (offset color for classic look)
    col2 = GREEN if i % 2 == 0 else BLUE
    triangle([(x0, botY), (x1, botY), ((x0+x1)//2, botY-tri_h)], col2)


def checker(cx, cy, d, color):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, int(cx-d/2), int(cy-d/2), int(d), int(d))
    c.fill.solid(); c.fill.fore_color.rgb = color
    c.line.color.rgb = WHITE; c.line.width = Pt(1)
    c.shadow.inherit = False
    return c


d = int(tri_base * 0.78)
# a few stacked checkers for flavour
for k in range(3):
    checker(quad_x(0)+tri_base/2, topY + d*0.6 + k*d*0.95, d, WHITE)
for k in range(2):
    checker(quad_x(11)+tri_base/2, botY - d*0.6 - k*d*0.95, d, RGBColor(0x22,0x2A,0x38))
# dice
for i, (dx, val) in enumerate([(0, None), (1, None)]):
    pass
die1 = rect(s, bar_x - Inches(0.0), pT + pH//2 - Inches(0.28), Inches(0.55), Inches(0.55),
            WHITE, radius=True)
# (kept minimal — board reads clearly)

# tagline bottom
tb, tf = box(s, Inches(0), SH - Inches(0.85), SW, Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER)
r = p.add_run(); r.text = "פרויקט בקורס מבני נתונים"
style(r, 16, MUTED, italic=True)


# =========================================================================
# SLIDE 2 — HISTORY
# =========================================================================
s = slide()
accent_header(s, "היסטוריית המשחק", 2)
tb, tf = box(s, Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.6))
para(tf, first=True, space_after=18)
r = tf.paragraphs[0].add_run()
r.text = "שש־בש הוא משחק לוח עתיק, אך נֶשבֶש הוא וריאציה ייחודית משלו."
style(r, 22, WHITE, bold=True)

p = para(tf, space_after=14)
r = p.add_run()
r.text = ("הגרסה הנוכחית של המשחק הומצאה לפני למעלה מעשור, על ידיי "
          "ועל ידי חבר טוב שלי – נֶש, שעל שמו נקרא המשחק.")
style(r, 19, MUTED)

p = para(tf, space_after=14)
r = p.add_run()
r.text = ("פיתחנו יחד מערכת חוקים ייחודית של “הטלות מיוחדות” המוסיפות "
          "דרמה, טקטיקה והפתעות מעבר לשש־בש הקלאסי.")
style(r, 19, MUTED)

p = para(tf, space_after=14)
r = p.add_run()
r.text = ("המשחק הופץ מפה לאוזן בין חברים וזכה להצלחה מסחררת – "
          "וכך נולד הרעיון להפוך אותו לאפליקציה דיגיטלית שכולם יוכלו לשחק בה.")
style(r, 19, MUTED)

# decorative quote chip
chip = rect(s, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), PANEL, radius=True)
tb, tf = box(s, Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True)
r = p.add_run(); r.text = "“ממשחק שולחן בין חברים – לאפליקציה.”"
style(r, 18, GOLD, italic=True, bold=True)


# =========================================================================
# SLIDE 3 — GENERAL RULES
# =========================================================================
s = slide()
accent_header(s, "חוקי המשחק הכלליים", 3)
# two columns
colW = Inches(5.85)
tb, tf = box(s, SW - Inches(0.8) - colW, Inches(1.95), colW, Inches(5.0))
bullet(tf, "מבנה הלוח", "מערך של 26 תאים; 0 ו־25 הם ה“בר”. ערך חיובי = לבן, שלילי = שחור.", first=True)
bullet(tf, "מערכת 2 קליקים", "קליק ראשון מסמן יעדים: כחול = ביניים, ירוק = יעד סופי.")
bullet(tf, "נגעת – נסעת", "לאחר הקליק השני המהלך סופי. אין ביטול.")
bullet(tf, "דאבל", "דאבל רגיל מעניק תור נוסף אוטומטי.")

tb, tf = box(s, Inches(0.8), Inches(1.95), colW, Inches(5.0))
bullet(tf, "הטלת פתיחה", "כל שחקן מטיל קובייה אחת; הגבוה מתחיל ומשחק את שתי הקוביות.", first=True)
bullet(tf, "3 דאבלים ברצף", "“היפוך השולחן” – התור מסתיים מיד ועובר ליריב.")
bullet(tf, "ניקוד", "פשוט = 1, מארס = 2, מארס טורקי = 3.")
bullet(tf, "מארס כוכב", "ניצחון אליפות מיידי!")
bullet(tf, "מבנה טורניר", "ראשון ל־3 סטים; כל סט ראשון ל־3 נקודות.")


# =========================================================================
# SLIDE 4 — SPECIAL ROLLS
# =========================================================================
s = slide()
accent_header(s, "ההטלות המיוחדות", 4)

cards = [
    ("1 : 2", "דילוג", "התור מסתיים מיד.", BLUE),
    ("4 : 5", "בחירת דאבל", "השחקן בוחר כל דאבל לשחק.", GREEN),
    ("6 : 5", "מכת הנֶש", "כל ה“בלוטים” של היריב עפים לבר + 2 מהלכים חופשיים.", GOLD),
    ("6 : 3", "החלפה", "לשחק 6:3 או לזרוק מחדש.", BLUE),
    ("5 : 2", "אחורה", "מזיזים 5 ו־2 (או 7) לאחור.", GREEN),
    ("4 : 3", "טריגר ידני", "מטילים קובייה אחת = מספר הצעדים אחורה.", BLUE),
    ("5 : 1", "ארבעה מהלכים", "מטילים קובייה; התוצאה d מזכה ב־4 מהלכים בערך d.", GREEN),
]
# 4 + 3 grid
cw, ch = Inches(2.85), Inches(2.2)
gx, gy = Inches(0.7), Inches(2.0)
gap = Inches(0.18)
positions = [(0,0),(1,0),(2,0),(3,0),(0,1),(1,1),(2,1)]
for (label, name, desc, col), (cx, cy) in zip(cards, positions):
    x = gx + cx*(cw+gap); y = gy + cy*(ch+gap)
    card = rect(s, x, y, cw, ch, PANEL, radius=True)
    # accent bar at top
    rect(s, x, y, cw, Inches(0.12), col)
    tbb, tff = box(s, x+Inches(0.1), y+Inches(0.22), cw-Inches(0.2), ch-Inches(0.3))
    p = para(tff, first=True, align=PP_ALIGN.CENTER, rtl=False, space_after=4)
    r = p.add_run(); r.text = label; style(r, 26, col, bold=True)
    p = para(tff, align=PP_ALIGN.CENTER, space_after=4)
    r = p.add_run(); r.text = name; style(r, 17, WHITE, bold=True)
    p = para(tff, align=PP_ALIGN.CENTER, space_after=0)
    r = p.add_run(); r.text = desc; style(r, 12.5, MUTED)


# =========================================================================
# SLIDE 5 — ARCHITECTURE
# =========================================================================
s = slide()
accent_header(s, "מבנה ואדריכלות המערכת", 5)

colW = Inches(5.85)
# RIGHT column: data structures (course relevance)
tb, tf = box(s, SW - Inches(0.8) - colW, Inches(1.95), colW, Inches(5.0))
p = para(tf, first=True, space_after=10)
r = p.add_run(); r.text = "מבני נתונים בליבה"; style(r, 20, GREEN, bold=True)
bullet(tf, "מערך number[26]", "ייצוג הלוח כולו במערך יחיד – הבר בקצוות.", size=17)
bullet(tf, "מנוע חישוב מהלכים", "אלגוריתם המחשב את כל המסלולים האפשריים מהקוביות.", size=17)
bullet(tf, "מכונת מצבים", "ניהול תור, הטלות מיוחדות ומעברי מצב.", size=17)
bullet(tf, "ניהול מצב – Zustand", "Store יחיד כמקור אמת לכל נתוני המשחק.", size=17)

# LEFT column: tech stack
tb, tf = box(s, Inches(0.8), Inches(1.95), colW, Inches(5.0))
p = para(tf, first=True, space_after=10)
r = p.add_run(); r.text = "טכנולוגיות וכלים"; style(r, 20, BLUE, bold=True)
bullet(tf, "React Native + Expo", "אפליקציה חוצת־פלטפורמות ב־TypeScript.", size=17, lead_color=BLUE)
bullet(tf, "Reanimated / Moti", "אנימציות אכילה, היפוך שולחן וקוביות.", size=17, lead_color=BLUE)
bullet(tf, "Firebase RTDB", "ריבוי משתתפים בארכיטקטורת Host־Authoritative.", size=17, lead_color=BLUE)
bullet(tf, "expo-linking + QR", "קישורי הצטרפות עמוקים וסריקת קוד.", size=17, lead_color=BLUE)
bullet(tf, "Vercel", "פריסת גרסת הווב לשחקנים ללא התקנה.", size=17, lead_color=BLUE)


# =========================================================================
# SLIDE 6 — AI DEVELOPMENT
# =========================================================================
s = slide()
accent_header(s, "פיתוח בעזרת כלי AI", 6)
tb, tf = box(s, Inches(0.8), Inches(1.85), Inches(11.7), Inches(0.7))
p = para(tf, first=True)
r = p.add_run()
r.text = "הפרויקט פותח כמעט כולו בשיתוף פעולה עם כלי בינה מלאכותית:"
style(r, 19, WHITE, bold=True)

ai = [
    ("Claude", "מנוע הפיתוח המרכזי – לוגיקת המשחק, החוקים, הארכיטקטורה והקוד.", BLUE),
    ("Gemini", "סיוע ברעיונות, מחקר, ניסוח חוקים ופתרון בעיות.", GREEN),
    ("THE BIG O", "הבוט הייעודי של הקורס – ליווי וייעוץ במבני נתונים ואלגוריתמים.", GOLD),
]
cw, ch = Inches(3.7), Inches(3.6)
gx = Inches(0.85); gy = Inches(2.75); gap = Inches(0.25)
for i, (name, desc, col) in enumerate(ai):
    x = gx + i*(cw+gap)
    card = rect(s, x, gy, cw, ch, PANEL, radius=True)
    card.line.color.rgb = col; card.line.width = Pt(1.5)
    rect(s, x, gy, cw, Inches(0.5), col)
    tbb, tff = box(s, x, gy+Inches(0.05), cw, Inches(0.45), anchor=MSO_ANCHOR.MIDDLE)
    p = para(tff, first=True, align=PP_ALIGN.CENTER, rtl=False)
    r = p.add_run(); r.text = name; style(r, 22, RGBColor(0x0F,0x1B,0x2D), bold=True)
    tbb, tff = box(s, x+Inches(0.25), gy+Inches(0.9), cw-Inches(0.5), ch-Inches(1.1),
                   anchor=MSO_ANCHOR.MIDDLE)
    p = para(tff, first=True, align=PP_ALIGN.CENTER)
    r = p.add_run(); r.text = desc; style(r, 16, MUTED)


# =========================================================================
# SLIDE 7 — FUTURE
# =========================================================================
s = slide()
accent_header(s, "המשך הפרויקט", 7)
tb, tf = box(s, Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.8))
bullet(tf, "שיפור הממשק והעיצוב", "ליטוש חוויית המשתמש, אנימציות ועיצוב לוח מוקפד יותר.", first=True, size=20, space_after=16)
bullet(tf, "תיקון באגים", "איתון ותיקון תקלות בניסוי וטעייה לאורך המון משחקים אמיתיים.", size=20, space_after=16)
bullet(tf, "אפשרויות שיתוף נוספות", "דרכים קלות יותר להזמין חברים ולהצטרף למשחק.", size=20, space_after=16)
bullet(tf, "תכונות חדשות", "מצבי משחק נוספים, סטטיסטיקות ושיפורי ריבוי משתתפים.", size=20, space_after=16)
bullet(tf, "נגישות וזמינות", "הרחבת גרסת הווב כך שכל אחד יוכל לשחק מכל מכשיר.", size=20, space_after=16)


# =========================================================================
# SLIDE 8 — QR
# =========================================================================
s = slide()
URL = "https://nesh-besh.vercel.app/"
qr = qrcode.QRCode(border=2, box_size=14,
                   error_correction=qrcode.constants.ERROR_CORRECT_M)
qr.add_data(URL); qr.make(fit=True)
img = qr.make_image(fill_color="#0F1B2D", back_color="white")
qr_path = "presentation/qr_neshbesh.png"
img.save(qr_path)

tb, tf = box(s, Inches(0), Inches(0.7), SW, Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER, rtl=False)
r = p.add_run(); r.text = "Nesh"; style(r, 44, BLUE, bold=True)
r = p.add_run(); r.text = "Besh"; style(r, 44, GREEN, bold=True)

tb, tf = box(s, Inches(0), Inches(1.7), SW, Inches(0.7), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER)
r = p.add_run(); r.text = "סרקו את הקוד והצטרפו למשחק!"
style(r, 26, WHITE, bold=True)

# QR with white frame
qsz = Inches(3.4)
fx = (SW - qsz)//2
fy = Inches(2.55)
frame = rect(s, fx-Inches(0.2), fy-Inches(0.2), qsz+Inches(0.4), qsz+Inches(0.4),
             WHITE, radius=True)
s.shapes.add_picture(qr_path, fx, fy, qsz, qsz)

tb, tf = box(s, Inches(0), fy+qsz+Inches(0.25), SW, Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER, rtl=False)
r = p.add_run(); r.text = URL; style(r, 20, GOLD, bold=True)

tb, tf = box(s, Inches(0), SH - Inches(0.8), SW, Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, align=PP_ALIGN.CENTER)
r = p.add_run(); r.text = "תודה רבה!"; style(r, 18, MUTED, italic=True)


out = "presentation/NeshBesh.pptx"
prs.save(out)
print("Saved:", out)
